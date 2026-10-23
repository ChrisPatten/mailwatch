"""Attachment and fetched-document parsers. Text out; nothing is ever executed or unpacked."""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
from pathlib import Path
from typing import Any

log = logging.getLogger("mailwatch.attachments")

PDF_PAGES = 30
SHEET_ROWS = 200
TEXT_CAP = 40000
IMG_MIN_BYTES = 20_000
IMG_MIN_PIXELS = 200 * 200

EXT = {"pdf": ".pdf", "ics": ".ics", "docx": ".docx", "xlsx": ".xlsx", "csv": ".csv", "txt": ".txt",
       "html": ".html", "image": ".img", "other": ".bin", "json": ".json"}


def kind_of(ctype: str, name: str) -> str:
    n = name.lower()
    c = (ctype or "").lower()
    if c == "application/pdf" or n.endswith(".pdf"):
        return "pdf"
    if c == "text/calendar" or n.endswith(".ics"):
        return "ics"
    if n.endswith(".docx") or "wordprocessingml" in c:
        return "docx"
    if n.endswith(".xlsx") or "spreadsheetml" in c:
        return "xlsx"
    if n.endswith(".csv") or c == "text/csv":
        return "csv"
    if c.startswith("text/html") or n.endswith((".htm", ".html")):
        return "html"
    if c == "application/json" or n.endswith(".json"):
        return "json"
    if c.startswith("text/") or n.endswith((".txt", ".md")):
        return "txt"
    if c.startswith("image/"):
        return "image"
    return "other"


# --- parsers ---------------------------------------------------------------

def pdf_text(data: bytes) -> str:
    import pdfplumber
    out: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for i, page in enumerate(pdf.pages[:PDF_PAGES]):
            t = page.extract_text() or ""
            if len(t.strip()) < 20:
                t = _ocr_pdf_page(page) or t
            out.append(f"[page {i + 1}]\n{t.strip()}")
    return "\n\n".join(out)


def _ocr_pdf_page(page: Any) -> str:
    try:
        img = page.to_image(resolution=200).original
        return ocr_image(img)
    except Exception as e:
        log.debug("pdf page OCR unavailable: %s", e)
        return ""


def ocr_image(img: Any) -> str:
    """Apple Vision via ocrmac (accurate recognition + language correction). macOS only."""
    try:
        from ocrmac import ocrmac  # macOS only
    except Exception:
        return ""
    try:
        res = ocrmac.OCR(img, recognition_level="accurate", language_preference=["en-US"]).recognize()
        return "\n".join(r[0] for r in res)
    except Exception as e:
        log.debug("ocr failed: %s", e)
        return ""


def image_text(data: bytes, inline: bool) -> tuple[str, str]:
    if len(data) < IMG_MIN_BYTES:
        return "", "skipped: small image (logo or pixel)"
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        if img.width * img.height < IMG_MIN_PIXELS:
            return "", "skipped: small image"
        t = ocr_image(img)
        return t, "ocr" if t else "ocr unavailable or no text"
    except Exception as e:
        return "", f"unreadable image: {e}"


def docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def xlsx_text(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[sheet {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= SHEET_ROWS:
                out.append("…")
                break
            out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def csv_text(data: bytes) -> str:
    s = data.decode("utf-8", "replace")
    rows = list(csv.reader(io.StringIO(s)))[:SHEET_ROWS]
    return "\n".join(" | ".join(r) for r in rows)


def parse_ics(data: bytes) -> list[dict[str, Any]]:
    try:
        from icalendar import Calendar
        cal = Calendar.from_ical(data)
    except Exception:
        return []
    method = str(cal.get("METHOD", "") or "")
    evs = []
    for ev in cal.walk("VEVENT"):
        def g(k: str) -> str:
            v = ev.get(k)
            if v is None:
                return ""
            dt = getattr(v, "dt", None)
            return dt.isoformat() if dt is not None and hasattr(dt, "isoformat") else str(v)
        org = str(ev.get("ORGANIZER", "") or "").replace("mailto:", "").replace("MAILTO:", "")
        evs.append({"title": g("SUMMARY")[:200], "start": g("DTSTART"), "end": g("DTEND"),
                    "location": g("LOCATION")[:200], "organizer": org[:120], "method": method,
                    "status": g("STATUS"), "source": "ics"})
    return evs


def ics_text(evs: list[dict[str, Any]]) -> str:
    return "\n".join(f"{e['method'] or 'EVENT'}: {e['title']} | {e['start']} - {e['end']} | {e['location']} | "
                     f"organizer {e['organizer']}" for e in evs)


def html_doc_text(data: bytes) -> str:
    from .normalize import _LinkBook, clean_chars, collapse, html_to_text
    try:
        import trafilatura
        t = trafilatura.extract(data.decode("utf-8", "replace"), include_links=False, include_tables=True) or ""
    except Exception:
        t = ""
    if not t.strip():
        t, _ = html_to_text(data.decode("utf-8", "replace"), _LinkBook())
    t, _ = clean_chars(t)
    return collapse(t)


def extract_text(kind: str, data: bytes, inline: bool = False) -> tuple[str, str]:
    """Returns (text, note)."""
    try:
        if kind == "pdf":
            t = pdf_text(data)
            return t, f"PDF, {t.count('[page ')} pages read"
        if kind == "docx":
            return docx_text(data), "Word document"
        if kind == "xlsx":
            return xlsx_text(data), "spreadsheet"
        if kind == "csv":
            return csv_text(data), "CSV"
        if kind in ("txt", "json"):
            return data.decode("utf-8", "replace"), kind
        if kind == "html":
            return html_doc_text(data), "web page"
        if kind == "image":
            return image_text(data, inline)
        if kind == "ics":
            return ics_text(parse_ics(data)), "calendar invite"
    except Exception as e:
        return "", f"parse failed: {type(e).__name__}"
    return "", "not parsed (name and size only)"


def materialize(msg_id: str, items: list[dict[str, Any]], art_dir: Path | str) -> list[dict[str, Any]]:
    """Write bytes + extracted text under art_dir. Returns artifact rows (no bytes)."""
    art_dir = Path(art_dir)
    art_dir.mkdir(parents=True, exist_ok=True)
    out = []
    for i, it in enumerate(items):
        data: bytes = it["data"]
        aid = "a_" + hashlib.sha256(msg_id.encode() + data).hexdigest()[:10]
        kind = it["kind"]
        path = art_dir / f"{aid}{EXT.get(kind, '.bin')}"
        path.write_bytes(data)
        if "text" in it:
            text, note = it["text"], "calendar invite"
        else:
            text, note = extract_text(kind, data, it.get("inline", False))
        text = re.sub(r"\n{3,}", "\n\n", text or "")[:TEXT_CAP]
        tpath = ""
        if text.strip():
            tp = art_dir / f"{aid}.txt"
            tp.write_text(text)
            tpath = str(tp)
        out.append({"id": aid, "msg_id": msg_id, "kind": kind, "source": it.get("source", "attachment"),
                    "name": str(it.get("name", ""))[:200], "path": str(path), "text_path": tpath,
                    "size": len(data), "note": note})
    return out
